import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.normalize import normalize_docx
from app.profile import ConversionProfile
from app.validate import validate_docx


def _profile():
    return ConversionProfile(preserve=["legal-structure", "tables", "footnotes"])


def _source(path: Path):
    document = Document()
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = "18 . 8 . 90 Official Journal of the European Communities No L 224/47"
    footer = section.footer.paragraphs[0]
    footer.text = "18.8.90 Official Journal of the European Communities No L 224/47"
    document.add_paragraph("First legal paragraph.")
    second = document.add_paragraph("Bold and italic legal text.")
    second.runs[0].bold = True
    second.runs[0].italic = True
    document.add_paragraph("Second legal paragraph")
    document.add_paragraph("ANNEX I")
    annex_body = document.add_paragraph("Annex body remains a paragraph.")
    annex_body.paragraph_format.page_break_before = True
    document.add_paragraph("Ordinary mention of Annex in a sentence remains inline.")
    document.save(path)


def test_translation_ready_normalization_rules(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    _source(source)
    normalize_docx(source, output, _profile())
    report = validate_docx(output)
    assert report["editableCharacters"] > 60
    document = Document(output)
    paragraphs = list(document.paragraphs)
    texts = [paragraph.text for paragraph in paragraphs]
    assert not any("Official Journal" in text for text in texts)
    assert texts.count("First legal paragraph.") == 1
    assert texts.count("Second legal paragraph") == 1
    assert "Annex body remains a paragraph." in texts
    assert all(paragraph.style.name == "Normal" for paragraph in paragraphs)
    assert all(run.font.name == "Times New Roman" and run.font.size == Pt(12)
               for paragraph in paragraphs for run in paragraph.runs if run.text)
    styled = next(paragraph for paragraph in paragraphs if paragraph.text.startswith("Bold"))
    assert styled.runs[0].bold is True
    assert styled.runs[0].italic is True
    # Exactly one explicit page break is inserted before the Annex paragraph;
    # source page-break-before and all ordinary source breaks are removed.
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert xml.count('w:type="page"') == 1
    assert 'w:type="column"' not in xml
    assert '<w:sectPr' in xml
    assert xml.count('<w:sectPr') == 1


def test_annex_mentions_inside_sentence_do_not_break(tmp_path):
    source, output = tmp_path / "source.docx", tmp_path / "output.docx"
    document = Document()
    document.add_paragraph("This sentence mentions Annex I as a reference.")
    document.add_paragraph("ANNEX II")
    document.save(source)
    normalize_docx(source, output, _profile())
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert xml.count('w:type="page"') == 1


def test_legal_structural_alignment_and_oj_cleanup(tmp_path):
    source, output = tmp_path / "legal-source.docx", tmp_path / "legal-output.docx"
    document = Document()
    document.add_paragraph("18. 8. 90 Official Journal of the European Communities No L 224/47")
    document.add_paragraph("COUNCIL DIRECTIVE")
    document.add_paragraph("of 26 June 1990")
    document.add_paragraph("on animal health conditions governing the movement of equidae")
    document.add_paragraph("(90/426/EEC)")
    document.add_paragraph("Whereas animals must be protected;")
    document.add_paragraph("CHAPTER I")
    document.add_paragraph("General provisions")
    document.add_paragraph("Article 1")
    document.add_paragraph("Scope")
    document.add_paragraph("This Directive lays down ordinary legal provisions.")
    document.add_paragraph("Short body text")
    document.add_paragraph("(1) OJ No C 327, 30.12.1989, p. 1. (2) OJ No C 149, 18.6.1990, p. 2.")
    document.add_paragraph("No L 224/44 Official Journal of the European Communities 18. 8. 90")
    document.save(source)

    normalize_docx(source, output, _profile())
    result = Document(output)
    by_text = {paragraph.text: paragraph for paragraph in result.paragraphs}
    assert not any("Official Journal of the European Communities" in text for text in by_text)
    for text in (
        "COUNCIL DIRECTIVE", "of 26 June 1990",
        "on animal health conditions governing the movement of equidae", "(90/426/EEC)",
        "CHAPTER I", "General provisions", "Article 1", "Scope",
        "(1) OJ No C 327, 30.12.1989, p. 1. (2) OJ No C 149, 18.6.1990, p. 2.",
    ):
        assert by_text[text].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert by_text[text].style.name == "Normal"
    for text in ("Whereas animals must be protected;", "This Directive lays down ordinary legal provisions.", "Short body text"):
        assert by_text[text].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    reference = by_text["(1) OJ No C 327, 30.12.1989, p. 1. (2) OJ No C 149, 18.6.1990, p. 2."]
    assert all(run.font.size == Pt(9) for run in reference.runs if run.text)
    assert reference.paragraph_format.space_before == Pt(0)
    assert reference.paragraph_format.space_after == Pt(0)


@pytest.mark.parametrize("identifier", [
    "(90/426/EEC)",
    "(EU) 2024/1234",
    "(EC) No 123/2004",
    "(EEC) No 90/426",
])
def test_standalone_legal_act_identifiers_are_centered(tmp_path, identifier):
    source, output = tmp_path / "identifier-source.docx", tmp_path / "identifier-output.docx"
    document = Document()
    document.add_paragraph(identifier)
    document.save(source)

    normalize_docx(source, output, _profile())

    paragraph = next(paragraph for paragraph in Document(output).paragraphs if paragraph.text == identifier)
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.style.name == "Normal"
    assert all(run.font.name == "Times New Roman" and run.font.size == Pt(12) for run in paragraph.runs if run.text)


def test_legal_act_identifier_inside_body_sentence_is_not_centered(tmp_path):
    source, output = tmp_path / "body-reference-source.docx", tmp_path / "body-reference-output.docx"
    text = "The conditions in (EU) 2024/1234 continue to apply."
    document = Document()
    document.add_paragraph(text)
    document.save(source)

    normalize_docx(source, output, _profile())

    paragraph = next(paragraph for paragraph in Document(output).paragraphs if paragraph.text == text)
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_empty_page_boundary_paragraphs_are_removed(tmp_path):
    source, output = tmp_path / "boundary-source.docx", tmp_path / "boundary-output.docx"
    document = Document()
    document.add_paragraph("Paragraph before source page boundary.")
    boundary = document.add_paragraph()
    boundary.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph()
    document.add_paragraph("Paragraph after source page boundary.")
    document.save(source)
    normalize_docx(source, output, _profile())
    result = Document(output)
    assert [paragraph.text for paragraph in result.paragraphs] == [
        "Paragraph before source page boundary.",
        "Paragraph after source page boundary.",
    ]
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:type="page"' not in xml

