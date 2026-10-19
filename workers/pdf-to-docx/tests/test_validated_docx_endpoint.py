import hashlib
import json
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app import main

FIXTURE = Path(__file__).parent / "fixtures" / "32006L0114.pdf"
GENERATION = "10000000-0000-4000-8000-000000000001"
SOURCE = "20000000-0000-4000-8000-000000000002"
MODEL = "30000000-0000-4000-8000-000000000003"


def _model(monkeypatch):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "render-token")
    source = FIXTURE.read_bytes()
    response = TestClient(main.app).post("/reconstruct/model", headers={"Authorization": "Bearer render-token"},
        files={"document": ("fixture.pdf", source, "application/pdf")}, data={"generation_id": GENERATION,
        "source_artifact_id": SOURCE, "source_sha256": hashlib.sha256(source).hexdigest(),
        "config_sha256": hashlib.sha256(b"{}").hexdigest()})
    assert response.status_code == 200
    return response.content, hashlib.sha256(source).hexdigest()


def _render(model, source_sha):
    return TestClient(main.app).post("/render/validated-docx", headers={"Authorization": "Bearer render-token"},
        files={"model": ("model.json", model, "application/json")}, data={"model_generation_id": GENERATION,
        "model_artifact_id": MODEL, "model_sha256": hashlib.sha256(model).hexdigest(),
        "source_artifact_id": SOURCE, "source_sha256": source_sha})


def test_validated_model_renders_native_structurally_valid_docx(monkeypatch, tmp_path):
    model, source_sha = _model(monkeypatch)
    response = _render(model, source_sha)
    assert response.status_code == 200
    assert response.headers["x-docx-validation"] == "PASS"
    target = tmp_path / "validated.docx"
    target.write_bytes(response.content)
    document = Document(target)
    assert len(document.tables) == 3
    assert sum(len(table.rows) for table in document.tables) == 48
    assert sum(len(row.cells) for table in document.tables for row in table.rows) == 100
    assert all(len(cell.paragraphs) == 1 for table in document.tables for row in table.rows for cell in row.cells)


def test_corrupt_model_bytes_fail_closed(monkeypatch):
    model, source_sha = _model(monkeypatch)
    corrupt = model + b" "
    response = TestClient(main.app).post("/render/validated-docx", headers={"Authorization": "Bearer render-token"},
        files={"model": ("model.json", corrupt, "application/json")}, data={"model_generation_id": GENERATION,
        "model_artifact_id": MODEL, "model_sha256": hashlib.sha256(model).hexdigest(),
        "source_artifact_id": SOURCE, "source_sha256": source_sha})
    assert response.status_code == 422
    assert "SHA-256" in response.json()["error"]["details"]["reason"]


def test_schema_and_provenance_mismatches_fail_before_render(monkeypatch):
    model, source_sha = _model(monkeypatch)
    payload = json.loads(model)
    payload["schemaVersion"] = "future"
    changed = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    assert _render(changed, source_sha).status_code == 422
    payload = json.loads(model)
    inline = next(node for node in _walk(payload["document"]) if node["type"] == "InlineSpan")
    inline["source"][0]["span_ids"] = []
    changed = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
    assert _render(changed, source_sha).status_code == 422


def test_malformed_docx_output_is_rejected(monkeypatch):
    model, source_sha = _model(monkeypatch)
    monkeypatch.setattr(main, "render_docx_phase2b", lambda _document, target: target.write_bytes(b"not-a-docx"))
    response = _render(model, source_sha)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "docx_serialization_validation_failed"


def _walk(node):
    yield node
    for child in node["children"]:
        yield from _walk(child)
