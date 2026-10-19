from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

import pytest
from lxml import etree

from app.canonical import canonicalize_docx
from app.reconstruction.contracts import Box, Classification, Fragment, Node, SourceRef, StyleEvidence
from app.reconstruction.core import assert_text_integrity, build_typed_document, iter_content_nodes, render_docx_phase2b, _is_ambiguous_line_break_hyphen, _is_safe_line_break_hyphen, docx_visible_text
from app.reconstruction.footnote_structure import bind_footnotes

FIXTURE = Path(__file__).parent / "fixtures" / "32006L0114.pdf"
W = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _node(identity: str, text: str, page: int, y0: float, y1: float, kind: str = "Paragraph") -> Node:
    ref = SourceRef(page, 1, [f"p{page}-{identity}"], [f"p{page}-{identity}-s"], [Box(page, 40, y0, 500, y1)], 0, 0)
    return Node(identity, kind, text, [ref], StyleEvidence(["Times"], [9], False, False, None, 40), Classification("deterministic", 1, ["fixture"]))


def test_32006l0114_binds_all_source_footnotes_and_writes_native_ooxml(tmp_path: Path):
    document, decisions, _ = build_typed_document(FIXTURE, "32006L0114")
    bodies = [node for node in document.children if node.type == "FootnoteBody"]
    assert [len(node.attributes["footnoteEntries"]) for node in bodies] == [4, 2, 1]
    assert len(document.attributes["footnotes"]) == 7
    assert document.attributes["footnoteNumbering"] == {
        "scheme": "page-local-explicit",
        "pages": {"1": ["1", "2", "3", "4"], "2": ["1", "2"], "3": ["1"]},
    }
    assert len([item for item in decisions if item.kind == "footnote-link"]) == 7
    target = tmp_path / "footnotes.docx"; render_docx_phase2b(document, target)
    assert assert_text_integrity(document, target)["footnotes"] == 7
    with zipfile.ZipFile(target) as archive:
        assert "word/footnotes.xml" in archive.namelist()
        body = etree.fromstring(archive.read("word/document.xml")); notes = etree.fromstring(archive.read("word/footnotes.xml")); styles = etree.fromstring(archive.read("word/styles.xml"))
        assert len(body.xpath(".//w:footnoteReference", namespaces=W)) == 7
        actual_notes = notes.xpath("./w:footnote[not(@w:type)]", namespaces=W)
        assert len(actual_notes) == 7
        references = body.xpath(".//w:footnoteReference", namespaces=W)
        assert all(ref.get(f"{{{W['w']}}}customMarkFollows") == "1" for ref in references)
        assert all(ref.getparent().find(f"{{{W['w']}}}t") is not None for ref in references)
        assert all(ref.getparent().xpath("./w:rPr/w:rStyle[@w:val='FootnoteReference']", namespaces=W) for ref in references)
        assert [ref.getparent().xpath("string(./w:t)", namespaces=W) for ref in references] == ["(1)", "(2)", "(3)", "(4)", "(1)", "(2)", "(1)"]
        assert styles.xpath("./w:style[@w:type='character' and @w:styleId='FootnoteReference']/w:rPr/w:vertAlign[@w:val='superscript']", namespaces=W)
        assert styles.xpath("./w:style[@w:type='paragraph' and @w:styleId='FootnoteText']", namespaces=W)
        assert not notes.xpath("./w:footnote[not(@w:type)]//w:footnoteRef", namespaces=W)
        assert notes.xpath("./w:footnote[@w:type='separator' and @w:id='-1']", namespaces=W)
        assert notes.xpath("./w:footnote[@w:type='continuationSeparator' and @w:id='0']", namespaces=W)
        settings = etree.fromstring(archive.read("word/settings.xml"))
        assert settings.xpath("./w:footnotePr/w:footnote[@w:id='-1']", namespaces=W)
        assert settings.xpath("./w:footnotePr/w:footnote[@w:id='0']", namespaces=W)
        expected_markers = ["1", "2", "3", "4", "1", "2", "1"]
        assert [note.xpath("./w:p[1]/w:r[1]/w:t/text()", namespaces=W)[0] for note in actual_notes] == expected_markers
        body_text = " ".join(body.xpath(".//w:body//w:t/text()", namespaces=W))
        assert "Opinion of 26 October 2006" not in body_text


def test_canonical_projection_exposes_native_footnotes_and_reference_codes(tmp_path: Path):
    document, _, _ = build_typed_document(FIXTURE, "32006L0114")
    target = tmp_path / "footnotes.docx"; render_docx_phase2b(document, target)
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    payload, metrics = canonicalize_docx(target, validated_docx_artifact_id="fixture", validated_docx_sha256=digest)
    envelope = json.loads(payload)
    assert metrics == {"nodeCount": 373, "translationUnitCount": 222, "inlineCodeCount": 31, "normalizedTokens": 3257}
    assert len([unit for unit in envelope["translationUnits"] if unit["unit_type"] == "footnote"]) == 7
    assert len([code for code in envelope["inlineCodes"] if code["code_type"] == "footnote_reference"]) == 7


def test_missing_and_ambiguous_anchors_fail_closed():
    body = _node("body", "(1) OJ L 1, p. 1.", 1, 760, 790, "ListItem")
    document = _node("document", "", 1, 0, 800, "Document"); document.children = [body]
    with pytest.raises(ValueError, match="0 safe anchors"):
        bind_footnotes(document, {1: 800}, [])
    ambiguous_body = _node("ambiguous-body", "(1) OJ L 1, p. 1.", 1, 760, 790, "ListItem")
    document.children = [
        _node("a", "Text (1).", 1, 100, 120),
        _node("b", "Other (1).", 1, 150, 170),
        ambiguous_body,
    ]
    with pytest.raises(ValueError, match="2 safe anchors"):
        bind_footnotes(document, {1: 800}, [])


def test_page_boundary_does_not_supply_an_anchor_from_another_page():
    anchor = _node("anchor", "Text (1).", 1, 100, 120)
    body = _node("body", "(1) OJ L 1, p. 1.", 2, 760, 790, "Footnote")
    document = _node("document", "", 1, 0, 800, "Document"); document.children = [anchor, body]
    with pytest.raises(ValueError, match="page 2 has 0 safe anchors"):
        bind_footnotes(document, {1: 800, 2: 800}, [])


def test_legal_subdivision_is_not_confused_with_footnote_call():
    anchor = _node("anchor", "Article 8(1) applies; citation (1).", 1, 100, 120)
    body = _node("body", "(1) OJ L 1, p. 1.", 1, 760, 790, "Footnote")
    document = _node("document", "", 1, 0, 800, "Document"); document.children = [anchor, body]
    records = bind_footnotes(document, {1: 800}, [])
    assert records[0]["anchorOffset"] == anchor.exact_text.rindex("(1)")


def test_numbering_restarts_and_multiple_notes_are_modeled_per_source_page():
    nodes = [
        _node("a1", "First (1).", 1, 100, 120),
        _node("a2", "Second (2).", 1, 140, 160),
        _node("f1", "(1) Note one. (2) Note two.", 1, 700, 790, "Footnote"),
        _node("a3", "Third (1).", 2, 100, 120),
        _node("f2", "(1) Note three.", 2, 740, 790, "Footnote"),
    ]
    document = _node("document", "", 1, 0, 800, "Document"); document.children = nodes
    bind_footnotes(document, {1: 800, 2: 800}, [])
    assert document.attributes["footnoteNumbering"]["pages"] == {"1": ["1", "2"], "2": ["1"]}


def test_oj_page_furniture_is_excluded_with_provenance_and_no_blank_paragraphs(tmp_path: Path):
    document, decisions, diagnostics = build_typed_document(FIXTURE, "32006L0114")
    furniture = [item for item in decisions if item.kind == "page-furniture"]
    assert len(furniture) == 28
    assert len({ref.page for item in furniture for ref in item.source}) == 7
    assert all(item.ambiguity_status == "resolved" and item.outcome == "excluded" for item in furniture)
    assert not any(node.exact_text.strip() == "Official Journal of the European Union" for node in document.children if node.exact_text)
    target = tmp_path / "no-furniture.docx"
    render_docx_phase2b(document, target)
    from docx import Document
    output = Document(target)
    texts = [paragraph.text.strip() for paragraph in output.paragraphs]
    assert not any(text in {"EN", "Official Journal of the European Union"} for text in texts)
    assert not any(text in {"27.12.2006", "L 376/21", "L 376/22", "L 376/23", "L 376/24", "L 376/25", "L 376/26", "L 376/27"} for text in texts)
    blank_paragraphs = [paragraph for paragraph in output.paragraphs if not paragraph.text.strip()]
    assert len(blank_paragraphs) == 0
    assert not any(paragraph._p.xpath(".//w:br[@w:type='page']") for paragraph in output.paragraphs)


def test_renderer_uses_explicit_source_typography_without_word_heading_or_list_styles(tmp_path: Path):
    document, _, _ = build_typed_document(FIXTURE, "32006L0114")
    target = tmp_path / "typography.docx"
    render_docx_phase2b(document, target)
    with zipfile.ZipFile(target) as archive:
        document_xml = etree.fromstring(archive.read("word/document.xml"))
        styles_xml = etree.fromstring(archive.read("word/styles.xml"))
        assert not document_xml.xpath(".//w:numPr", namespaces=W)
        assert not document_xml.xpath(".//w:pStyle[starts-with(@w:val, 'Heading') or @w:val='Title' or contains(@w:val, 'List')]", namespaces=W)
        assert not document_xml.xpath(".//w:rStyle[starts-with(@w:val, 'Heading') or @w:val='Title' or contains(@w:val, 'List')]", namespaces=W)
        assert not document_xml.xpath(".//w:color[not(@w:val='000000')]", namespaces=W)
        assert not styles_xml.xpath(".//w:style[@w:default='1' and @w:type='paragraph'][starts-with(@w:styleId, 'Heading') or @w:styleId='Title']", namespaces=W)


def test_renderer_applies_standard_times_new_roman_spacing_and_alignment(tmp_path: Path):
    document, _, _ = build_typed_document(FIXTURE, "32006L0114")
    target = tmp_path / "standard-format.docx"
    render_docx_phase2b(document, target)
    with zipfile.ZipFile(target) as archive:
        body = etree.fromstring(archive.read("word/document.xml"))
        notes = etree.fromstring(archive.read("word/footnotes.xml"))
        paragraphs = body.xpath(".//w:p[.//w:t]", namespaces=W)
        assert paragraphs
        assert all(p.xpath("./w:pPr/w:spacing[@w:before='120' and @w:after='120' and @w:line='240']", namespaces=W) for p in paragraphs)
        assert all(p.xpath("./w:pPr/w:ind[@w:left='567' and @w:hanging='567']", namespaces=W) or p.xpath("./w:pPr/w:ind[not(@w:left) or @w:left='0'][not(@w:right) or @w:right='0'][not(@w:firstLine) or @w:firstLine='0']", namespaces=W) or not p.xpath("./w:pPr/w:ind", namespaces=W) for p in paragraphs)
        runs = body.xpath(".//w:r[.//w:t]", namespaces=W)
        assert all(run.xpath("./w:rPr/w:rFonts[@w:ascii='Times New Roman' and @w:hAnsi='Times New Roman']", namespaces=W) for run in runs)
        assert all(run.xpath("./w:rPr/w:sz[@w:val='24' or @w:val='22' or @w:val='20']", namespaces=W) for run in runs)
        assert all(run.xpath("./w:rPr/w:color[@w:val='000000']", namespaces=W) for run in runs)
        note_runs = notes.xpath("./w:footnote[not(@w:type)]//w:r[.//w:t]", namespaces=W)
        assert all(run.xpath("./w:rPr/w:rFonts[@w:ascii='Times New Roman']", namespaces=W) for run in note_runs)


def test_recitals_are_single_tabbed_hanging_paragraphs(tmp_path: Path):
    document, _, _ = build_typed_document(FIXTURE, "32006L0114")
    recitals = [node for node in document.children if node.type == "Recital"]
    assert [node.attributes["recitalNumber"] for node in recitals] == list(range(1, 22))
    target = tmp_path / "recitals.docx"
    render_docx_phase2b(document, target)
    with zipfile.ZipFile(target) as archive:
        body = etree.fromstring(archive.read("word/document.xml"))
        recital_paragraphs = [p for p in body.xpath(".//w:body/w:p", namespaces=W)
                             if p.xpath("./w:pPr/w:tabs/w:tab[@w:pos='567']", namespaces=W)]
        assert len(recital_paragraphs) == 21
        assert all(p.xpath("./w:pPr/w:ind[@w:left='567' and @w:hanging='567']", namespaces=W) for p in recital_paragraphs)
        assert all(len(p.xpath(".//w:r/w:tab", namespaces=W)) == 1 for p in recital_paragraphs)
        assert not body.xpath(".//w:numPr", namespaces=W)
    assert assert_text_integrity(document, target)["status"] == "PASS"


def test_articles_are_centered_and_annexes_are_the_only_body_page_breaks(tmp_path: Path):
    document, _, _ = build_typed_document(FIXTURE, "32006L0114")
    target = tmp_path / "structure-format.docx"
    render_docx_phase2b(document, target)
    with zipfile.ZipFile(target) as archive:
        body = etree.fromstring(archive.read("word/document.xml"))
        articles = [p for p in body.xpath(".//w:body/w:p", namespaces=W) if "".join(p.xpath(".//w:t/text()", namespaces=W)).strip().startswith("Article ")]
        assert articles
        assert all(p.xpath("./w:pPr/w:jc[@w:val='center']", namespaces=W) for p in articles)
        assert not body.xpath(".//w:br[@w:type='page']", namespaces=W)
        annexes = [p for p in body.xpath(".//w:body/w:p", namespaces=W) if p.xpath("./w:pPr/w:pageBreakBefore", namespaces=W)]
        assert len(annexes) == 2
        assert not body.xpath(".//w:pPr/w:pageBreakBefore[preceding-sibling::w:pageBreakBefore]", namespaces=W)
        notes = etree.fromstring(archive.read("word/footnotes.xml"))
        note_paragraphs = notes.xpath("./w:footnote[not(@w:type)]/w:p", namespaces=W)
        assert note_paragraphs
        assert all(p.xpath("./w:pPr/w:spacing[@w:before='0' and @w:after='0' and @w:line='240']", namespaces=W) for p in note_paragraphs)
        assert all(p.xpath(".//w:sz[@w:val='20']", namespaces=W) for p in note_paragraphs)
        assert all(p.xpath("./w:pPr/w:pStyle[@w:val='FootnoteText']", namespaces=W) for p in note_paragraphs)
        assert all(p.xpath("./w:pPr/w:ind[@w:left='340' and @w:hanging='340']", namespaces=W) for p in note_paragraphs)
        assert all(p.xpath("./w:pPr/w:tabs/w:tab[@w:val='left' and @w:pos='340']", namespaces=W) for p in note_paragraphs)
        assert all(len(p.xpath("./w:r[1]/w:tab", namespaces=W)) == 1 for p in note_paragraphs)
        assert all(len(p.xpath(".//w:footnoteReference", namespaces=W)) == 0 for p in note_paragraphs)
        assert all(p.xpath("./w:r[1]/w:rPr/w:vertAlign[@w:val='superscript']", namespaces=W) for p in note_paragraphs)
        assert all(p.xpath("./w:r[1]/w:rPr/w:rStyle[@w:val='FootnoteReference']", namespaces=W) for p in note_paragraphs)
        assert all(not ("(" in "".join(p.xpath("./w:r[1]/w:t/text()", namespaces=W)) or ")" in "".join(p.xpath("./w:r[1]/w:t/text()", namespaces=W))) for p in note_paragraphs)
        assert not notes.xpath(".//w:pBdr", namespaces=W)


def test_line_break_hyphenation_requires_geometry_and_preserves_compounds():
    left = Fragment(1, 2, "Mem-", Box(1, 10, 100, 100, 112), ["p1-b2-l4"], [], [], [], False, False, [])
    right = Fragment(1, 2, "ber", Box(1, 10, 114, 100, 126), ["p1-b2-l5"], [], [], [], False, False, [])
    assert _is_safe_line_break_hyphen(left.text, right.text, left, right)
    compound_left = Fragment(1, 2, "well-", Box(1, 10, 100, 100, 112), ["p1-b2-l4"], [], [], [], False, False, [])
    assert not _is_safe_line_break_hyphen(compound_left.text, right.text, compound_left, right)
    assert _is_ambiguous_line_break_hyphen(compound_left.text, right.text)
    same_line = Fragment(1, 2, "Mem-", Box(1, 10, 100, 100, 112), ["p1-b2-l4"], [], [], [], False, False, [])
    same_line_next = Fragment(1, 2, "ber", Box(1, 10, 100, 100, 112), ["p1-b2-l4"], [], [], [], False, False, [])
    assert not _is_safe_line_break_hyphen(same_line.text, same_line_next.text, same_line, same_line_next)


def test_known_split_words_are_reconstructed_in_body_and_footnotes(tmp_path: Path):
    document, _, diagnostics = build_typed_document(FIXTURE, "32006L0114")
    source = "\n".join(node.exact_text for node in iter_content_nodes(document))
    assert "con- cerning" not in source
    assert "inter- ests" not in source
    assert "Commu- nity" not in source
    assert "Jour- nal" not in source
    assert any(item["left"].endswith("con-") and item["right"].startswith("cerning") for item in diagnostics["extraction"]["dehyphenation"])
    assert any(item["left"].endswith("Jour-") and item["right"].startswith("nal") for item in diagnostics["extraction"]["dehyphenation"])
    target = tmp_path / "dehyphenated.docx"
    render_docx_phase2b(document, target)
    visible = docx_visible_text(target)
    assert "con- cerning" not in visible
    assert "inter- ests" not in visible
    assert "Commu- nity" not in visible
    assert "Jour- nal" not in visible
    assert assert_text_integrity(document, target)["status"] == "PASS"
