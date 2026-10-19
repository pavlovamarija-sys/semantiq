import json
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt
from fastapi.testclient import TestClient

from app import main
from app.backend import Pdf2DocxBackend
from app.diagnostics import PdfDiagnostics
from app.errors import WorkerError
from app.validate import validate_docx
from scripts.docx_validator import validate_docx as validate_application_docx


PROFILE = {
    "output": "docx",
    "editableText": True,
    "preserve": ["legal-structure", "tables", "footnotes"],
    "page": {"columns": 1, "marginsCm": {"top": 2.54, "bottom": 2.54, "left": 3.17, "right": 3.17}},
    "body": {"alignment": "justified", "beforePt": 6, "afterPt": 6, "lineSpacing": "single"},
    "footnotes": {"fontSizePt": 9, "beforePt": 0, "afterPt": 0, "lineSpacing": "single"},
    "literalLegalLabels": True,
}


def _fake_conversion(_self, _source, target):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "No L 224/44 Official Journal of the European Communities 18 . 8 . 90"
    document.add_paragraph("18 . 8 . 90 Official Journal of the European Communities No L 224/47")
    document.add_paragraph("COUNCIL DIRECTIVE")
    document.add_paragraph("of 26 June 1990")
    document.add_paragraph("on animal health conditions governing the movement of equidae")
    document.add_paragraph("(90/426/EEC)")
    document.add_heading("Article 1", level=1)
    document.add_paragraph("Scope")
    styled = document.add_paragraph("This Regulation establishes a legally binding test paragraph (1).")
    styled.runs[0].bold = True
    styled.runs[0].italic = True
    document.add_paragraph("(a) first legal subdivision")
    document.add_paragraph("(1) OJ No C 327, 30.12.1989, p. 1.")
    boundary = document.add_paragraph()
    boundary.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("ANNEX I")
    document.add_paragraph("Annex title")
    document.add_paragraph("Annex body")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Source"
    table.cell(0, 1).text = "Value"
    document.save(target)
    return "native-pdf2docx", PdfDiagnostics((), 80, 12, 80, "native_text")


def test_endpoint_returns_valid_editable_docx(monkeypatch, tmp_path):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    monkeypatch.setenv("PDF_TO_DOCX_CONVERTER_MODE", "legacy")
    monkeypatch.setattr(Pdf2DocxBackend, "convert", _fake_conversion)
    response = TestClient(main.app).post(
        "/convert/pdf-to-docx",
        headers={"Authorization": "Bearer test-token"},
        files={"document": ("legal.pdf", b"%PDF-1.7\nfixture", "application/pdf")},
        data={"profile": json.dumps(PROFILE)},
    )
    assert response.status_code == 200
    assert response.headers["x-ooxml-validated"] == "true"
    assert response.headers["x-normalization-validated"] == "true"
    assert response.headers["x-converter-engine"].startswith("pdf2docx/")
    assert response.headers["x-conversion-path"] == "native-pdf2docx"
    output = tmp_path / "result.docx"
    output.write_bytes(response.content)
    report = validate_docx(output)
    assert report["editableCharacters"] > 20
    with zipfile.ZipFile(BytesIO(response.content)) as archive:
        xml = archive.read("word/document.xml")
        assert b"Article 1" in xml
        assert b"first legal subdivision" in xml
        assert b'w:left="1797"' in xml
        assert b'w:num="1"' in xml
        decoded = xml.decode("utf-8")
        assert "Official Journal of the European Communities" not in decoded
        assert decoded.count('w:type="page"') == 1
        assert 'w:type="column"' not in decoded
        assert decoded.count("<w:sectPr") == 1

    # This is the exact final byte stream returned for download, not the raw
    # converter output or an in-memory pre-save document.
    final = Document(BytesIO(response.content))
    by_text = {paragraph.text: paragraph for paragraph in final.paragraphs}
    assert all(paragraph.style.name == "Normal" for paragraph in final.paragraphs)
    assert all(
        run.font.name == "Times New Roman"
        and run.font.size == (Pt(9) if paragraph.text.startswith("(1) OJ") else Pt(12))
        for paragraph in final.paragraphs for run in paragraph.runs if run.text
    )
    for text in ("COUNCIL DIRECTIVE", "of 26 June 1990", "on animal health conditions governing the movement of equidae",
                 "(90/426/EEC)", "Article 1", "Scope", "ANNEX I", "Annex title"):
        assert by_text[text].alignment == WD_ALIGN_PARAGRAPH.CENTER
    automatic_reference_paragraph = next(
        paragraph for paragraph in final.paragraphs
        if paragraph.text.startswith("This Regulation establishes a legally binding test paragraph")
    )
    assert automatic_reference_paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    for text in ("(a) first legal subdivision", "Annex body"):
        assert by_text[text].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    preserved = automatic_reference_paragraph.runs[0]
    assert preserved.bold is True and preserved.italic is True


def test_analysis_endpoint_uses_canonical_pymupdf_extraction(monkeypatch, tmp_path):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    source = tmp_path / "source.pdf"
    import pymupdf
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 100), "Article 1 Canonical source text.")
    document.save(source)
    response = TestClient(main.app).post(
        "/analyze/pdf",
        headers={"Authorization": "Bearer test-token"},
        files={"document": ("legal.pdf", source.read_bytes(), "application/pdf")},
    )
    assert response.status_code == 200
    report = response.json()
    assert report["integrityValidated"] is True
    assert report["sourceText"].strip() == "Article 1 Canonical source text."
    assert report["charactersWithoutSpaces"] == len("Article1Canonicalsourcetext.")
    assert report["charactersWithSpaces"] == len("Article 1 Canonical source text.")
    assert report["diagnostics"]["classification"] == "native_text"
    assert report["countError"] is None


def test_rejects_non_pdf(monkeypatch):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    response = TestClient(main.app).post(
        "/convert/pdf-to-docx",
        headers={"Authorization": "Bearer test-token"},
        files={"document": ("not.pdf", b"plain text", "application/pdf")},
        data={"profile": json.dumps(PROFILE)},
    )
    assert response.status_code == 415
    assert response.json()["error"]["code"] == "invalid_pdf"


def test_requires_authentication(monkeypatch):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    response = TestClient(main.app).post(
        "/convert/pdf-to-docx",
        files={"document": ("legal.pdf", b"%PDF-1.7\nfixture", "application/pdf")},
        data={"profile": json.dumps(PROFILE)},
    )
    assert response.status_code == 401
    assert response.json()["error"]["code"] == "unauthorized"


def test_worker_failure_preserves_status_code_message_and_details(monkeypatch):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    monkeypatch.setenv("PDF_TO_DOCX_CONVERTER_MODE", "legacy")

    def fail_conversion(*_args, **_kwargs):
        raise WorkerError("footnotes_need_review", "Footnote matching needs review; source text was preserved.", details={"marker": "1", "candidateCount": 2})

    monkeypatch.setattr(main.ConversionService, "convert", fail_conversion)
    response = TestClient(main.app).post(
        "/convert/pdf-to-docx",
        headers={"Authorization": "Bearer test-token"},
        files={"document": ("legal.pdf", b"%PDF-1.7\nfixture", "application/pdf")},
        data={"profile": json.dumps(PROFILE)},
    )
    assert response.status_code == 422
    assert response.json() == {"error": {
        "code": "footnotes_need_review",
        "message": "Footnote matching needs review; source text was preserved.",
        "details": {"marker": "1", "candidateCount": 2},
    }}


def test_typed_converter_is_default_and_preserves_endpoint_contract(monkeypatch, tmp_path):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    monkeypatch.delenv("PDF_TO_DOCX_CONVERTER_MODE", raising=False)
    source = Path(__file__).parent / "fixtures" / "32025D2586.pdf"
    response = TestClient(main.app).post(
        "/convert/pdf-to-docx",
        headers={"Authorization": "Bearer test-token"},
        files={"document": (source.name, source.read_bytes(), "application/pdf")},
        data={"profile": json.dumps(PROFILE)},
    )
    assert response.status_code == 200, response.text
    assert response.headers["x-converter-engine"] == main.TYPED_CONVERTER_ENGINE
    assert response.headers["x-conversion-path"] == main.TYPED_CONVERSION_PATH
    assert response.headers["x-ooxml-validated"] == "true"
    assert response.headers["x-normalization-validated"] == "true"
    assert response.headers["x-reading-order-validated"] == "true"
    target = tmp_path / "typed.docx"
    target.write_bytes(response.content)
    document = Document(target)
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 5
    assert len(document.tables[0].columns) == 5
    assert validate_docx(target)["footnotes"] == 1
    application_report = validate_application_docx(target)
    assert application_report["tables"] == 1
    assert application_report["footnotes"] == 1


def test_typed_failure_never_falls_back_to_legacy(monkeypatch):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    monkeypatch.setenv("PDF_TO_DOCX_CONVERTER_MODE", "typed")
    legacy_called = False

    def legacy_conversion(*_args, **_kwargs):
        nonlocal legacy_called
        legacy_called = True

    monkeypatch.setattr(main.ConversionService, "convert", legacy_conversion)
    monkeypatch.setattr(main, "_validate_typed_conversion", lambda *_args: (_ for _ in ()).throw(ValueError("gate failed")))
    response = TestClient(main.app).post(
        "/convert/pdf-to-docx",
        headers={"Authorization": "Bearer test-token"},
        files={"document": ("legal.pdf", b"%PDF-1.7\nfixture", "application/pdf")},
        data={"profile": json.dumps(PROFILE)},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "typed_reconstruction_validation_failed"
    assert legacy_called is False


def test_invalid_converter_mode_fails_closed(monkeypatch):
    monkeypatch.setenv("PDF_TO_DOCX_SERVICE_TOKEN", "test-token")
    monkeypatch.setenv("PDF_TO_DOCX_CONVERTER_MODE", "automatic")
    response = TestClient(main.app).post(
        "/convert/pdf-to-docx",
        headers={"Authorization": "Bearer test-token"},
        files={"document": ("legal.pdf", b"%PDF-1.7\nfixture", "application/pdf")},
        data={"profile": json.dumps(PROFILE)},
    )
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "configuration_error"
