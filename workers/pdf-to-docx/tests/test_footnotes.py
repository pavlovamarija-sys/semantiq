import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from app.errors import WorkerError
from app.footnotes import _split_notes, create_word_footnotes
from app.diagnostics import analyze_pdf
from app.profile import ConversionProfile
from app.service import ConversionService


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
NS = {"w": W_NS}


def test_spaced_footnote_marker_keeps_the_complete_oj_citation():
    assert _split_notes("( 1 ) OJ No C 149, 18.6.1990, p. 2.") == [
        ("1", "OJ No C 149, 18.6.1990, p. 2."),
    ]


def _xml(archive: zipfile.ZipFile, name: str):
    return etree.fromstring(archive.read(name))


def test_ambiguous_marker_keeps_source_text_for_review(tmp_path):
    path = tmp_path / "ambiguous.docx"
    document = Document()
    document.add_paragraph("First unrelated reference (2).")
    document.add_paragraph("The authoritative source marker is also (2).")
    document.add_paragraph("(2) OJ No C 149, 18.6.1990, p. 2.")
    document.save(path)

    with pytest.raises(WorkerError, match="needs review") as failure:
        create_word_footnotes(path)

    assert failure.value.code == "footnotes_need_review"
    assert failure.value.details["candidateCount"] == 2
    assert Document(path).paragraphs[-1].text == "(2) OJ No C 149, 18.6.1990, p. 2."
    with zipfile.ZipFile(path) as archive:
        assert "word/footnotes.xml" not in archive.namelist()


def test_90_426_eec_footnotes_are_real_word_footnotes(tmp_path):
    source = Path(__file__).parent / "fixtures" / "31990L0426.pdf"
    output = tmp_path / "31990L0426.docx"
    source_analysis = analyze_pdf(source)
    assert source_analysis.diagnostics.classification == "native_text"
    assert source_analysis.charactersWithoutSpaces == 36154
    assert source_analysis.charactersWithSpaces == 63888
    assert "COUNCIL DIRECTIVE" in source_analysis.sourceText
    report = ConversionService().convert(
        source,
        output,
        ConversionProfile(preserve=["legal-structure", "tables", "footnotes"]),
    )

    # The requested page-one example has three notes; the complete act has
    # three further genuine OJ source footnotes on later pages.
    assert report["footnotes"] == 6
    with zipfile.ZipFile(output) as archive:
        document = _xml(archive, "word/document.xml")
        footnotes = _xml(archive, "word/footnotes.xml")
        relationships = _xml(archive, "word/_rels/document.xml.rels")
        references = document.xpath(".//w:footnoteReference", namespaces=NS)
        notes = footnotes.xpath("w:footnote[number(@w:id) >= 0]", namespaces=NS)
        visible_text = "".join(
            node.text or ""
            for part in (document, footnotes)
            for node in part.xpath(".//w:t", namespaces=NS)
        )
        assert report["charactersWithoutSpaces"] == sum(not character.isspace() for character in visible_text)
        assert report["charactersWithSpaces"] == len(visible_text)
        assert len(references) == len(notes) == 6
        assert [node.get(f"{{{W_NS}}}id") for node in references[:3]] == ["1", "2", "3"]
        assert all(node.get(f"{{{W_NS}}}customMarkFollows") is None for node in references)
        assert all(note.get(f"{{{W_NS}}}customMarkFollows") is None for note in notes)
        assert all(len(note.xpath(".//w:footnoteRef", namespaces=NS)) == 1 for note in notes)
        assert all(
            note.xpath("string(.//w:footnoteRef/parent::w:r/w:rPr/w:rStyle/@w:val)", namespaces=NS) == "FootnoteReference"
            for note in notes
        )
        assert [" ".join("".join(note.xpath(".//w:t/text()", namespaces=NS)).split()) for note in notes[:3]] == [
            "OJ No C 327, 30. 12. 1989, p. 61 .",
            "OJ No C 149 , 18 . 6 . 1990.",
            "OJ No C 62, 12. 3 . 1990, p. 46 .",
        ]
        assert all("(*)" not in "".join(note.xpath(".//w:t/text()", namespaces=NS)) for note in notes)
        body_paragraphs = [
            "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
            for paragraph in document.xpath(".//w:body//w:p", namespaces=NS)
        ]
        body_text = " ".join(body_paragraphs)
        assert "OJ No C 327" not in body_text
        assert "OJ No C 149" not in body_text
        assert "OJ No C 62" not in body_text
        page_one_calls = [
            paragraph for paragraph in document.xpath(".//w:body//w:p", namespaces=NS)
            if paragraph.xpath(".//w:footnoteReference[@w:id='1' or @w:id='2' or @w:id='3']", namespaces=NS)
        ]
        assert len(page_one_calls) == 3
        assert [
            paragraph.xpath("string(.//w:footnoteReference/@w:id)", namespaces=NS)
            for paragraph in page_one_calls
        ] == ["1", "2", "3"]
        assert "(*)" not in " ".join("".join(p.xpath(".//w:t/text()", namespaces=NS)) for p in page_one_calls)
        assert len(relationships.xpath(
            "*[local-name()='Relationship' and @Type=$relationship_type]",
            relationship_type=FOOTNOTES_REL,
        )) == 1
        for note in notes:
            paragraphs = note.xpath("./w:p", namespaces=NS)
            assert len(paragraphs) == 1
            assert "".join(paragraphs[0].xpath(".//w:t/text()", namespaces=NS)).strip()
            for paragraph in note.xpath(".//w:p", namespaces=NS):
                spacing = paragraph.find(f"{{{W_NS}}}pPr/{{{W_NS}}}spacing")
                indentation = paragraph.find(f"{{{W_NS}}}pPr/{{{W_NS}}}ind")
                assert spacing is not None
                assert spacing.get(f"{{{W_NS}}}before") == "0"
                assert spacing.get(f"{{{W_NS}}}after") == "0"
                assert spacing.get(f"{{{W_NS}}}line") == "240"
                assert indentation is not None
                assert all(indentation.get(f"{{{W_NS}}}{key}") == "0" for key in ("left", "right", "firstLine", "hanging"))
            for run in note.xpath(".//w:r", namespaces=NS):
                size = run.find(f"{{{W_NS}}}rPr/{{{W_NS}}}sz")
                fonts = run.find(f"{{{W_NS}}}rPr/{{{W_NS}}}rFonts")
                assert size is not None and size.get(f"{{{W_NS}}}val") == "18"
                assert fonts is not None and fonts.get(f"{{{W_NS}}}ascii") == "Times New Roman"
                if run.xpath(".//w:t", namespaces=NS):
                    assert run.find(f"{{{W_NS}}}rPr/{{{W_NS}}}color").get(f"{{{W_NS}}}val") == "000000"
                    assert run.find(f"{{{W_NS}}}rPr/{{{W_NS}}}rStyle") is None
