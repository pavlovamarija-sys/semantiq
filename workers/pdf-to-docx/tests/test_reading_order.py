from pathlib import Path

import pytest
from docx import Document

from app.errors import WorkerError
from app.reading_order import validate_reading_order
from app.reconstruct import reconstruct_native_docx


FIXTURE = Path(__file__).parent / "fixtures" / "31990L0426.pdf"


def _positions(texts: list[str], needles: list[str]) -> list[int]:
    return [next(i for i, text in enumerate(texts) if text.startswith(needle)) for needle in needles]


def test_90_426_eec_two_column_reading_order(tmp_path):
    output = tmp_path / "31990L0426.docx"
    reconstruct_native_docx(FIXTURE, output)
    report = validate_reading_order(output)
    texts = [" ".join(p.text.split()) for p in Document(output).paragraphs if p.text.strip()]

    transition = ["HAS ADOPTED THIS DIRECTIVE", "CHAPTER 1", "General provisions", "Article 1", "Article 2", "(d) 'equidae for slaughter'", "(i) 'temporary admission'", "CHAPTER II", "Rules for the movement of equidae", "Article 3"]
    assert _positions(texts, transition) == sorted(_positions(texts, transition))
    annexes = ["ANNEX A", "COMPULSORILY NOTIFIABLE DISEASES", "ANNEX B", "ANNEX C", "ANNEX D"]
    assert _positions(texts, annexes) == sorted(_positions(texts, annexes))
    assert report["readingOrderValidated"] is True
    assert not any("Official Journal of th" in text for text in texts)


def test_impossible_article_order_fails_before_ready(tmp_path):
    output = tmp_path / "corrupt.docx"
    document = Document()
    for text in ("HAS ADOPTED THIS DIRECTIVE:", "CHAPTER I", "Article 1", "Article 3", "Article 2"):
        document.add_paragraph(text)
    document.save(output)

    with pytest.raises(WorkerError, match="reading-order validation failed") as error:
        validate_reading_order(output)
    assert error.value.code == "invalid_reading_order"
