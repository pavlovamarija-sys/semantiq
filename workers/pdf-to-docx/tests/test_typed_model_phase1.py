from pathlib import Path

import pytest
from docx import Document

from spikes.typed_model_phase1 import (
    Box, Classification, DiagnosticRecord, Node, SourceRef, StyleEvidence, _ambiguity_diagnostic,
    _log_layout_diagnostics, _resolve_paragraph_boundaries, assert_source_model_integrity,
    assert_text_integrity, build_typed_document, inspect_docx_structure, render_docx, render_docx_phase2b,
    validate_structural_reading_order, classify_paragraph_boundary,
)


SPIKE_FIXTURE = Path(__file__).parent / "fixtures" / "32006L0114.pdf"


@pytest.mark.skipif(not SPIKE_FIXTURE.exists(), reason="Neutral Phase 1 fixture is generated/downloaded separately.")
def test_32006l0114_typed_tables_and_integrity(tmp_path):
    model, decisions, diagnostics = build_typed_document(SPIKE_FIXTURE)
    tables = [node for node in model.children if node.type == "Table"]
    assert [(table.attributes["rows"], table.attributes["columns"]) for table in tables] == [(3, 2), (4, 3), (41, 2)]
    output = tmp_path / "experimental.docx"
    render_docx(model, output)
    assert assert_text_integrity(model, output)["status"] == "PASS"
    assert assert_source_model_integrity(model, diagnostics["sourceReadingText"], diagnostics["sourceSpanIds"])["status"] == "PASS"
    structure = inspect_docx_structure(output)
    assert structure["article1InTableCell"] is True
    assert structure["annexIInTableCell"] is True
    assert validate_structural_reading_order(model)["status"] == "PASS"
    assert any(decision.kind == "merge" for decision in decisions)
    assert sum(1 for decision in decisions if decision.kind == "split") == 3
    inline = next(child for node in model.children if node.type == "Paragraph" for child in node.children if child.type == "InlineSpan")
    assert inline.exact_text
    assert inline.source[0].span_ids
    def walk(node):
        yield node
        for child in node.children:
            yield from walk(child)
    assert all(node.source for node in walk(model))
    assert isinstance(diagnostics["unresolved"], list)


def test_structural_validator_ignores_table_cells_and_rejects_bad_headings():
    empty_style = __import__("spikes.typed_model_phase1", fromlist=["StyleEvidence"]).StyleEvidence([], [], False, False, None, 0)
    classification = __import__("spikes.typed_model_phase1", fromlist=["Classification"]).Classification("deterministic", 1, [])
    headings = [Node(f"a{value}", "Heading", f"Article {value}", [], empty_style, classification, attributes={"kind": "article", "ordinal": value}) for value in (1, 2, 3)]
    cell = Node("cell", "TableCell", "Article 1", [], empty_style, classification)
    row = Node("row", "TableRow", "Article 1", [], empty_style, classification, [cell])
    table = Node("table", "Table", "Article 1", [], empty_style, classification, [row], {"rows": 1, "columns": 1})
    root = Node("root", "Document", "", [], empty_style, classification, headings + [table])
    assert validate_structural_reading_order(root)["articles"] == [1, 2, 3]
    root.children.insert(2, Node("bad", "Heading", "Article 5", [], empty_style, classification, attributes={"kind": "article", "ordinal": 5}))
    with pytest.raises(ValueError, match="Article sequence"):
        validate_structural_reading_order(root)


@pytest.mark.skipif(not SPIKE_FIXTURE.exists(), reason="Neutral Phase 1 fixture is generated/downloaded separately.")
def test_phase2b_uses_one_paragraph_with_line_breaks_per_multiline_cell(tmp_path):
    model, _, _ = build_typed_document(SPIKE_FIXTURE)
    output = tmp_path / "phase2b.docx"
    render_docx_phase2b(model, output)
    assert assert_text_integrity(model, output)["status"] == "PASS"
    document = Document(output)
    part_a = document.tables[0]
    for row in part_a.rows:
        cell = row.cells[0]
        assert len(cell.paragraphs) == 1
        assert "\n" in cell.text


def test_aligned_list_rejection_and_layout_ambiguity_use_validated_diagnostics():
    rejection = _ambiguity_diagnostic("table-rejection", "aligned-body-columns-not-table", [],
        "Independent body-column flow is stronger than table evidence.", ["List indentation dominates."], .94)
    ambiguity = _ambiguity_diagnostic("layout-ambiguity", "gutter-straddler", [],
        "Two geometry interpretations remain viable.", ["Candidate: separator.", "Candidate: wide paragraph."], .65)
    logged = _log_layout_diagnostics([rejection, ambiguity], 4, [])
    assert [item.kind for item in logged] == ["table-rejection", "layout-ambiguity"]
    assert all(item.ambiguity_status == "unresolved" for item in logged)
    assert logged[0].reason != logged[0].evidence[0]


def test_table_and_cross_page_decisions_share_the_diagnostic_contract():
    table = DiagnosticRecord("classification", "table-1", [], "Repeated row and column geometry identifies a table.", .97,
        "classified-as-table", ["Three stable column anchors."], "not-ambiguous")
    cross_page = DiagnosticRecord("cross-page-table", "logical-table-1", [], "Compatible page-edge segments share one logical identity.", .96,
        "linked", ["Column boundaries agree.", "Repeated header agrees."], "resolved")
    assert table.outcome == "classified-as-table"
    assert cross_page.kind == "cross-page-table"
    assert cross_page.evidence == ["Column boundaries agree.", "Repeated header agrees."]


def test_missing_required_diagnostic_fields_fail_at_validation_boundary():
    with pytest.raises(ValueError, match="missing required fields: reason"):
        DiagnosticRecord.from_mapping({"kind":"layout-ambiguity","node_id":"x","confidence":.5,"outcome":"ambiguous","evidence":[],"ambiguity_status":"unresolved"})
    with pytest.raises(ValueError, match="reason is required"):
        DiagnosticRecord("layout-ambiguity", "x", [], "", .5, "ambiguous", [], "unresolved")
    with pytest.raises(ValueError, match="evidence"):
        DiagnosticRecord("layout-ambiguity", "x", [], "A reason.", .5, "ambiguous", [""], "unresolved")


def _boundary_node(node_id, text, block, line, order, box, node_type="Paragraph"):
    source = SourceRef(1, block, [f"p1-b{block}-l{line}"], [f"p1-b{block}-l{line}-s0"], [Box(1, *box)], order, order)
    style = StyleEvidence(["LegalSerif"], [9.6], False, False, None, box[0])
    inline = Node(f"span-{node_id}", "InlineSpan", text, [source], style, Classification("deterministic", 1, []))
    return Node(node_id, node_type, text, [source], style, Classification("deterministic", .9, []), [inline])


def test_generalized_boundary_classifier_merges_and_preserves_all_provenance():
    left = _boundary_node("left", "shall apply", 7, 1, 10, (50, 100, 100, 110), "ListItem")
    right = _boundary_node("right", "to every case", 7, 2, 11, (50, 111, 120, 121))
    classification = classify_paragraph_boundary(left, right)
    assert classification.decision == "MERGE"
    merged, counts = _resolve_paragraph_boundaries([left, right], [])
    assert counts == {"MERGE": 1, "KEEP_BOUNDARY": 0, "ABSTAIN": 0}
    assert merged[0].exact_text == "shall apply to every case"
    assert [ref.line_ids[0] for ref in merged[0].source] == ["p1-b7-l1", "p1-b7-l2"]
    assert [child.source[0].span_ids[0] for child in merged[0].children] == ["p1-b7-l1-s0", "p1-b7-l2-s0"]


def test_generalized_boundary_classifier_keeps_running_header_and_can_abstain():
    header = _boundary_node("header", "30.12.2006", 9, 0, 20, (500, 50, 550, 60))
    body = _boundary_node("body", "continued body text", 3, 1, 21, (330, 100, 550, 110))
    assert classify_paragraph_boundary(header, body).decision == "KEEP_BOUNDARY"
    uncertain_left = _boundary_node("uncertain-left", "possible continuation", 10, 1, 30, (50, 200, 200, 210))
    uncertain_right = _boundary_node("uncertain-right", "without close geometry", 11, 0, 31, (80, 260, 250, 270))
    assert classify_paragraph_boundary(uncertain_left, uncertain_right).decision == "ABSTAIN"
