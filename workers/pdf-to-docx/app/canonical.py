from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path
from docx import Document
from lxml import etree

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
SCHEMA = "eu-acquis.canonical-document.v1"
EXTRACTOR = "docx-canonical-phase6c5-v2-unicode"
FOOTNOTE_EXTRACTOR = "docx-canonical-phase6c8-v3-footnotes"


def _stable_id(prefix: str, value: str) -> str:
    return f"{prefix}-{hashlib.sha256(value.encode('utf-8')).hexdigest()[:24]}"


def _paragraph_text(paragraph) -> tuple[str, list[dict[str, object]]]:
    text, codes = [], []
    for child in paragraph._p.iterchildren():
        tag = etree.QName(child).localname
        if tag == "r":
            for part in child.iterchildren():
                local = etree.QName(part).localname
                if local == "t": text.append(part.text or "")
                elif local == "tab": codes.append({"kind": "tab", "offset": len("".join(text))}); text.append("\t")
                elif local == "br": codes.append({"kind": "line_break", "offset": len("".join(text))}); text.append("\n")
                elif local == "footnoteReference": codes.append({"kind": "footnote_reference", "offset": len("".join(text)), "footnote_id": part.get(f"{{{NS['w']}}}id")})
                elif local in {"drawing", "object", "pict", "fldChar", "instrText"}:
                    raise ValueError(f"Unsupported protected OOXML element: {local}")
        elif tag == "hyperlink":
            for part in child.xpath(".//w:t/text()", namespaces=NS): text.append(part)
        elif tag not in {"pPr", "bookmarkStart", "bookmarkEnd", "proofErr", "commentRangeStart", "commentRangeEnd"}:
            raise ValueError(f"Unsupported paragraph OOXML element: {tag}")
    value = "".join(text)
    for code in codes:
        if code["kind"] == "footnote_reference":
            marker = re.match(r"\((?:\d+|\*)\)", value[int(code["offset"]):])
            if marker is None: raise ValueError("Footnote reference has no preserved visible marker.")
            code["literal"] = marker.group(0)
    return value, codes


def canonicalize_docx(path: Path, *, validated_docx_artifact_id: str, validated_docx_sha256: str) -> tuple[bytes, dict[str, object]]:
    actual = hashlib.sha256(path.read_bytes()).hexdigest()
    if actual != validated_docx_sha256: raise ValueError("DOCX bytes do not match the immutable artifact hash.")
    with zipfile.ZipFile(path) as package:
        if package.testzip() is not None: raise ValueError("DOCX ZIP integrity failed.")
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml", "word/_rels/document.xml.rels"}
        if not required.issubset(package.namelist()): raise ValueError("DOCX package is incomplete.")
        xml = etree.fromstring(package.read("word/document.xml"))
        if xml.xpath(".//w:altChunk", namespaces=NS): raise ValueError("External altChunk content is unsupported.")
        footnotes_by_id = {}
        if "word/footnotes.xml" in package.namelist():
            footnotes = etree.fromstring(package.read("word/footnotes.xml"))
            for note in footnotes.xpath("./w:footnote[not(@w:type)]", namespaces=NS):
                note_id = note.get(f"{{{NS['w']}}}id")
                parts = []
                for run in note.xpath(".//w:r", namespaces=NS):
                    for child in run.iterchildren():
                        local = etree.QName(child).localname
                        if local == "t" and child.text:
                            parts.append(child.text)
                        elif local == "tab":
                            parts.append("\t")
                        elif local == "br":
                            parts.append("\n")
                footnotes_by_id[note_id] = "".join(parts)
    doc = Document(path); nodes, units, codes = [], [], []
    ordinal = 0
    tu_ordinal = 0
    def add_paragraph(paragraph, parent: str, location: str):
        nonlocal ordinal, tu_ordinal
        text, inline = _paragraph_text(paragraph)
        if not text.strip():
            if paragraph._p.xpath(".//w:br[@w:type='page']"):
                nodes.append({"id": _stable_id("node", f"{validated_docx_artifact_id}|{location}"), "node_type": "page_break", "ordinal": ordinal, "parent_id": parent, "source_location": location, "source_text": ""})
                ordinal += 1
            return
        node_id = _stable_id("node", f"{validated_docx_artifact_id}|{location}")
        tu_id = _stable_id("tu", f"{validated_docx_artifact_id}|{location}|{text}")
        recital = re.match(r"^\((\d+)\)\t", text)
        kind = "recital" if recital else ("heading" if paragraph.style.name.startswith("Heading") else ("footnote" if "footnote" in paragraph.style.name.lower() else "paragraph"))
        node_record = {"id": node_id, "node_type": kind, "ordinal": ordinal, "parent_id": parent, "source_location": location, "source_text": text}
        if recital:
            node_record["recital_number"] = int(recital.group(1))
        nodes.append(node_record)
        units.append({"id": tu_id, "node_id": node_id, "ordinal": tu_ordinal, "unit_type": kind, "source_text": text, "source_hash": hashlib.sha256(text.encode()).hexdigest()})
        tu_ordinal += 1
        for index, item in enumerate(inline): codes.append({"id": _stable_id("code", f"{tu_id}|{index}|{item['kind']}|{item['offset']}"), "translation_unit_id": tu_id, "ordinal": index, "code_type": item["kind"], "source_offset": item["offset"], "literal": item.get("literal", "\n" if item["kind"] == "line_break" else "\t")})
        ordinal += 1
        for item in inline:
            if item["kind"] != "footnote_reference": continue
            note_id, note_text = str(item["footnote_id"]), footnotes_by_id.get(str(item["footnote_id"]))
            if not note_text: raise ValueError(f"Footnote reference {note_id} has no body.")
            footnote_node_id = _stable_id("node", f"{validated_docx_artifact_id}|footnote:{note_id}")
            footnote_tu_id = _stable_id("tu", f"{validated_docx_artifact_id}|footnote:{note_id}|{note_text}")
            nodes.append({"id": footnote_node_id, "node_type": "footnote", "ordinal": ordinal, "parent_id": node_id, "source_location": f"footnote:{note_id}", "source_text": note_text})
            units.append({"id": footnote_tu_id, "node_id": footnote_node_id, "ordinal": tu_ordinal, "unit_type": "footnote", "source_text": note_text, "source_hash": hashlib.sha256(note_text.encode()).hexdigest()})
            ordinal += 1; tu_ordinal += 1
    def add_table(table, ti: int):
        nonlocal ordinal
        table_id = _stable_id("node", f"{validated_docx_artifact_id}|table:{ti}")
        nodes.append({"id": table_id, "node_type": "table", "ordinal": ordinal, "parent_id": "document", "source_location": f"table:{ti}", "source_text": ""}); ordinal += 1
        for ri, row in enumerate(table.rows):
            row_id = _stable_id("node", f"{table_id}|row:{ri}"); nodes.append({"id": row_id, "node_type": "table_row", "ordinal": ordinal, "parent_id": table_id, "source_location": f"table:{ti}/row:{ri}", "source_text": ""}); ordinal += 1
            for ci, cell in enumerate(row.cells):
                cell_id = _stable_id("node", f"{row_id}|cell:{ci}"); nodes.append({"id": cell_id, "node_type": "table_cell", "ordinal": ordinal, "parent_id": row_id, "source_location": f"table:{ti}/row:{ri}/cell:{ci}", "source_text": cell.text}); ordinal += 1
                for pi, paragraph in enumerate(cell.paragraphs): add_paragraph(paragraph, cell_id, f"table:{ti}/row:{ri}/cell:{ci}/p:{pi}")
    body_table_index = 0
    body_paragraph_index = 0
    for child in doc.element.body.iterchildren():
        local = etree.QName(child).localname
        if local == "p":
            paragraph = next((item for item in doc.paragraphs if item._p is child), None)
            if paragraph is not None: add_paragraph(paragraph, "document", f"body/p:{body_paragraph_index}")
            body_paragraph_index += 1
        elif local == "tbl":
            table = next((item for item in doc.tables if item._tbl is child), None)
            if table is not None: add_table(table, body_table_index)
            body_table_index += 1
    if not units: raise ValueError("DOCX contains no translatable source units.")
    normalized = " ".join(" ".join(unit["source_text"].split()) for unit in units)
    if any("\ufffd" in str(item.get("source_text", "")) for item in nodes + units):
        raise ValueError("DOCX source text contains Unicode replacement characters.")
    extractor = FOOTNOTE_EXTRACTOR if footnotes_by_id else EXTRACTOR
    envelope = {"schemaVersion": SCHEMA, "extractorVersion": extractor, "validatedDocxArtifactId": validated_docx_artifact_id, "validatedDocxSha256": validated_docx_sha256, "normalizedSourceSha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(), "nodes": nodes, "translationUnits": units, "inlineCodes": codes, "metrics": {"nodeCount": len(nodes), "translationUnitCount": len(units), "inlineCodeCount": len(codes), "normalizedTokens": len(normalized.split())}}
    return json.dumps(envelope, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(), envelope["metrics"]
