from __future__ import annotations

import re
from dataclasses import dataclass, replace
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from .errors import WorkerError

HEADING = re.compile(r"^(article\s+\d+|chapter\s+[ivxlcdm1]+|section\s+\d+|annex(?:\s+[ivxlcdm]+)?|title\s+[ivxlcdm]+)$", re.I)
RECITAL = re.compile(r"^\(\d+\)\s")
FOOTNOTE = re.compile(r"^\s*[\(\[]?[0-9*!�»x]+[\)\]]?\s*OJ\s+No\b", re.I)
OJ_HEADER = re.compile(r"official\s+journal\s+of\s+th(?:e|�|é)\s+european\s+communit", re.I)
OJ_NUMBER = re.compile(r"\bNo\s+L\s*\d+\s*/\s*\d+\b", re.I)
OJ_DATE = re.compile(r"\b\d{1,2}\s*\.\s*\d{1,2}\s*\.\s*\d{2,4}\b")


@dataclass(frozen=True)
class TextBlock:
    page: int
    x0: float
    y0: float
    x1: float
    y1: float
    text: str
    size: float
    bold: bool
    italic: bool


def _line_block(page_number: int, line: dict) -> TextBlock | None:
    spans = line.get("spans", [])
    text = "".join(s.get("text", "") or "".join(c.get("c", "") for c in s.get("chars", [])) for s in spans).strip()
    if not text:
        return None
    fonts = [s.get("font", "").lower() for s in spans]
    x0, y0, x1, y1 = line["bbox"]
    return TextBlock(page_number, x0, y0, x1, y1, text, max((float(s.get("size", 10)) for s in spans), default=10), any("bold" in f for f in fonts), any("italic" in f or "oblique" in f for f in fonts))


def _same_paragraph(previous: TextBlock, current: TextBlock) -> bool:
    line_height = max(previous.y1 - previous.y0, current.y1 - current.y0, 1)
    return current.y0 - previous.y1 <= line_height * 1.15 and abs(current.x0 - previous.x0) <= 25


def _blocks_from_dict(page_number: int, data: dict) -> list[TextBlock]:
    """Build from lines so a malformed PDF block can never bridge columns."""
    result: list[TextBlock] = []
    for raw in data.get("blocks", []):
        if raw.get("type") != 0:
            continue
        current = None
        for line in raw.get("lines", []):
            item = _line_block(page_number, line)
            if item is None:
                continue
            if current is not None and _same_paragraph(current, item):
                current = replace(current, x0=min(current.x0, item.x0), y0=min(current.y0, item.y0), x1=max(current.x1, item.x1), y1=max(current.y1, item.y1), text=f"{current.text} {item.text}", size=max(current.size, item.size), bold=current.bold or item.bold, italic=current.italic or item.italic)
            else:
                if current is not None:
                    result.append(current)
                current = item
        if current is not None:
            result.append(current)
    return result


def _is_running_margin(block: TextBlock, page_height: float) -> bool:
    if FOOTNOTE.match(block.text):
        return False
    if block.y0 > page_height * 0.92:
        return True
    if block.y1 > page_height * 0.14:
        return False
    text = " ".join(block.text.split())
    return bool(OJ_HEADER.search(text) or OJ_NUMBER.search(text) or OJ_DATE.fullmatch(text))


def _column_layout(blocks: list[TextBlock], page_width: float) -> tuple[float, float] | None:
    midpoint = page_width / 2
    left = [b for b in blocks if b.x1 < midpoint - page_width * 0.015]
    right = [b for b in blocks if b.x0 > midpoint + page_width * 0.015]
    if len(left) < 3 or len(right) < 3:
        return None
    overlaps = sum(1 for a in left for b in right if min(a.y1, b.y1) >= max(a.y0, b.y0))
    return (max(b.x1 for b in left), min(b.x0 for b in right)) if overlaps >= 3 else None


def _reading_order(blocks: list[TextBlock], page_width: float, page_height: float) -> list[TextBlock]:
    """Return top full-width material, then whole left and right columns."""
    content = [b for b in blocks if not _is_running_margin(b, page_height)]
    layout = _column_layout(content, page_width)
    if layout is None:
        return sorted(content, key=lambda b: (b.y0, b.x0))
    gutter_left, gutter_right = layout
    left_seed = [b for b in content if b.x1 < page_width / 2]
    right_seed = [b for b in content if b.x0 > page_width / 2]
    paired_y = min((max(a.y0, b.y0) for a in left_seed for b in right_seed if min(a.y1, b.y1) >= max(a.y0, b.y0)), default=0)
    full = [b for b in content if b.x0 < gutter_left and b.x1 > gutter_right]
    top_full = sorted((b for b in content if b.y0 < paired_y - 2 and (b in full or abs((b.x0 + b.x1) / 2 - page_width / 2) < page_width * 0.17)), key=lambda b: (b.y0, b.x0))
    remaining = [b for b in content if b not in top_full]
    left = sorted((b for b in remaining if (b.x0 + b.x1) / 2 < page_width / 2), key=lambda b: (b.y0, b.x0))
    right = sorted((b for b in remaining if (b.x0 + b.x1) / 2 >= page_width / 2), key=lambda b: (b.y0, b.x0))
    notes = [b for b in left + right if FOOTNOTE.match(b.text)]
    return top_full + [b for b in left if b not in notes] + [b for b in right if b not in notes] + notes


def _can_continue(previous: TextBlock, current: TextBlock) -> bool:
    if HEADING.match(previous.text) or HEADING.match(current.text) or FOOTNOTE.match(previous.text) or FOOTNOTE.match(current.text):
        return False
    if previous.text.rstrip().endswith((".", ";", ":", "?", "!")):
        return False
    return bool(current.text and (current.text[0].islower() or current.text[0] in "'\""))


def _write_blocks(pages: list[tuple[float, float, list[TextBlock]]], target: Path, *, ocr: bool) -> None:
    document = Document()
    document.core_properties.comments = "OCR fallback used" if ocr else "Native geometry reconstruction used"
    sizes = sorted(b.size for _, _, blocks in pages for b in blocks)
    normal_size = sizes[len(sizes) // 2] if sizes else 10.0
    previous = previous_paragraph = None
    for page_index, (width, height, blocks) in enumerate(pages):
        ordered = _reading_order(blocks, width, height)
        if page_index and not (previous and ordered and _can_continue(previous, ordered[0])):
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            previous_paragraph = None
        for block in ordered:
            clean = " ".join(block.text.split())
            if not clean:
                continue
            if previous is not None and previous_paragraph is not None and _can_continue(previous, block):
                previous_paragraph.add_run(f" {clean}")
                previous = block
                continue
            paragraph = document.add_heading(level=2) if HEADING.match(clean) or block.bold and block.size >= normal_size * 1.08 else document.add_paragraph()
            run = paragraph.add_run(clean)
            run.font.size = Pt(max(8, min(block.size, 18)))
            run.bold, run.italic = block.bold, block.italic
            if RECITAL.match(clean):
                paragraph.paragraph_format.left_indent = Pt(12)
            previous, previous_paragraph = block, paragraph
    document.save(target)


def _extract_pages(source: Path, mode: str) -> list[tuple[float, float, list[TextBlock]]]:
    with pymupdf.open(source) as document:
        return [(page.rect.width, page.rect.height, _blocks_from_dict(number, page.get_text(mode, sort=False))) for number, page in enumerate(document, 1)]


def reconstruct_raw_glyph_docx(source: Path, target: Path) -> None:
    pages = _extract_pages(source, "rawdict")
    if not any(blocks for _, _, blocks in pages):
        raise WorkerError("native_reconstruction_failed", "Raw PDF glyphs could not be reconstructed as editable text.")
    _write_blocks(pages, target, ocr=False)


def reconstruct_native_docx(source: Path, target: Path) -> None:
    pages = _extract_pages(source, "dict")
    if not any(blocks for _, _, blocks in pages):
        raise WorkerError("native_reconstruction_failed", "Native PDF text could not be reconstructed as editable text.")
    _write_blocks(pages, target, ocr=False)


def ocr_to_docx(source: Path, target: Path, language: str = "eng", dpi: int = 300) -> None:
    pages = []
    try:
        with pymupdf.open(source) as document:
            for number, page in enumerate(document, 1):
                textpage = page.get_textpage_ocr(language=language, dpi=dpi, full=True)
                pages.append((page.rect.width, page.rect.height, _blocks_from_dict(number, page.get_text("dict", textpage=textpage, sort=False))))
    except Exception as error:
        raise WorkerError("ocr_failed", "OCR fallback could not extract editable text.", details={"reason": str(error)}) from error
    if not any(blocks for _, _, blocks in pages):
        raise WorkerError("ocr_no_text", "OCR completed but found no editable text.")
    _write_blocks(pages, target, ocr=True)
