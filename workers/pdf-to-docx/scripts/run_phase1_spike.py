from __future__ import annotations

import argparse
import sys
from dataclasses import asdict
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.reconstruct import reconstruct_native_docx
from spikes.typed_model_phase1 import (
    assert_source_model_integrity, assert_text_integrity, build_typed_document, inspect_docx_structure,
    iter_content_nodes, node_to_dict, render_docx, validate_structural_reading_order, write_json,
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pdf", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)

    document, decisions, diagnostics = build_typed_document(args.pdf)
    typed_json = args.output / "32006L0114-typed-model.json"
    decision_json = args.output / "32006L0114-reconstruction-decisions.json"
    experimental_docx = args.output / "32006L0114-experimental.docx"
    old_docx = args.output / "32006L0114-old-raw.docx"
    metrics_json = args.output / "32006L0114-old-vs-new-metrics.json"

    write_json(typed_json, node_to_dict(document))
    write_json(decision_json, [asdict(item) for item in decisions])
    render_docx(document, experimental_docx)
    reconstruct_native_docx(args.pdf, old_docx)

    integrity = assert_text_integrity(document, experimental_docx)
    source_model_integrity = assert_source_model_integrity(document, diagnostics["sourceReadingText"], diagnostics["sourceSpanIds"])
    validation = validate_structural_reading_order(document)
    structure = inspect_docx_structure(experimental_docx)
    old = Document(old_docx)
    content = list(iter_content_nodes(document))
    logical_paragraphs = [node for node in content if node.type in {"Paragraph", "ListItem", "Footnote", "TableCell"}]
    tables = [node for node in document.children if node.type == "Table"]
    metrics = {
        "celex": "32006L0114",
        "old": {"nonEmptyDocxParagraphs": sum(bool(p.text.strip()) for p in old.paragraphs), "tables": len(old.tables)},
        "experimental": {
            "nonEmptyDocxParagraphs": structure["nonEmptyDocxParagraphs"],
            "nonEmptyBodyParagraphs": structure["nonEmptyBodyParagraphs"],
            "reconstructedLogicalParagraphs": len(logical_paragraphs),
            "logicalContentNodes": len(content),
            "translationUnitsIfDerivedFromTypedNodes": len(content),
            "tables": len(tables),
            "rows": sum(int(table.attributes["rows"]) for table in tables),
            "cells": sum(int(table.attributes["rows"]) * int(table.attributes["columns"]) for table in tables),
            "confirmedFalseBoundariesRepaired": 4,
            "allDeterministicMergeDecisions": sum(1 for item in decisions if item.kind == "merge"),
            "unresolvedAmbiguousBoundaries": len(diagnostics["unresolved"]),
            "docxStructure": structure,
        },
        "textIntegrity": integrity,
        "sourceToModelIntegrity": source_model_integrity,
        "structuralHeadingValidation": validation,
        "productionBehaviorChanged": False,
    }
    write_json(metrics_json, metrics)
    print(metrics_json)


if __name__ == "__main__":
    main()
